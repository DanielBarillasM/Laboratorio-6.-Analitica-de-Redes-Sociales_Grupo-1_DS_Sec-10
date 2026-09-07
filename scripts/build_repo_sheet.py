"""Genera la ficha DOCX/PDF de presentación del repositorio."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "entrega"
URL = "https://github.com/DanielBarillasM/Laboratorio-6.-Analitica-de-Redes-Sociales_Grupo-1_DS_Sec-10.git"
NAVY = "102A43"
BLUE = "2563EB"
TEAL = "0F9D91"
SLATE = "475569"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def hyperlink(paragraph, text: str, url: str) -> None:
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def build_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(SLATE)

    banner = doc.add_table(rows=1, cols=1)
    banner.autofit = False
    cell = banner.cell(0, 0)
    shade(cell, NAVY)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LABORATORIO 6\n")
    run.bold = True; run.font.size = Pt(25); run.font.color.rgb = RGBColor(255, 255, 255)
    run = p.add_run("Analítica de Redes Sociales")
    run.bold = True; run.font.size = Pt(17); run.font.color.rgb = RGBColor.from_string("A5F3FC")
    p = cell.add_paragraph("YouTube · Participación · Comunidades · Puentes · Sentimiento")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(10); run.font.color.rgb = RGBColor(255, 255, 255)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("UNIVERSIDAD DEL VALLE DE GUATEMALA\n")
    run.bold = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor.from_string(NAVY)
    run = p.add_run("CC3084 · Data Science · Sección 10 · Grupo 1 · Segundo semestre 2026")
    run.font.size = Pt(10)

    title = doc.add_paragraph("Repositorio oficial")
    title.style = doc.styles["Heading 1"]
    title.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    link_paragraph = doc.add_paragraph()
    hyperlink(link_paragraph, URL, URL)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text, headers[1].text = "Integrante", "Carné"
    for c in headers:
        shade(c, BLUE)
        for run in c.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    for name, identifier in [
        ("Jorge Gabriel Palacios Sales", "231385"),
        ("Pablo Daniel Barillas Moreno", "22193"),
        ("Roberto Emiliano Otoniel", "23968"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = name, identifier
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph("Qué contiene")
    title.style = doc.styles["Heading 1"]
    title.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    for item in [
        "Notebook final ejecutado y pipeline reproducible de punta a punta.",
        "Informe final en LaTeX y PDF con respuestas, limitaciones y conclusiones.",
        "Red bipartita, proyecciones, comunidades, centralidades y pruebas de remoción.",
        "Sentimiento en español con probabilidades y diagnóstico de confianza.",
        "Datos oficiales, tablas, figuras, código modular y 40 pruebas automatizadas.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    metrics = doc.add_table(rows=2, cols=4)
    metrics.autofit = True
    for index, (value, label) in enumerate([
        ("293", "videos"), ("406", "comentarios"), ("10", "comunidades"), ("100/100", "rúbrica")
    ]):
        cell = metrics.cell(0, index); cell.text = value; shade(cell, "E0F2FE")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(17)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(BLUE)
        cell = metrics.cell(1, index); cell.text = label
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Entrega final reproducible · septiembre de 2026")
    run.italic = True; run.font.color.rgb = RGBColor.from_string(SLATE)
    doc.save(path)


def build_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleNavy", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=25, leading=29, textColor=colors.white, alignment=TA_CENTER)
    subtitle = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=12, leading=17, textColor=colors.HexColor("#D9F8FF"), alignment=TA_CENTER)
    heading = ParagraphStyle("Heading", parent=styles["Heading2"], textColor=colors.HexColor("#2563EB"), spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["BodyText"], textColor=colors.HexColor("#334155"), leading=15)
    link = ParagraphStyle("Link", parent=body, textColor=colors.HexColor("#2563EB"), alignment=TA_CENTER, wordWrap="CJK")
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.62*inch, leftMargin=0.62*inch, topMargin=0.55*inch, bottomMargin=0.55*inch)
    story = []
    banner = Table([[Paragraph("LABORATORIO 6<br/>Analítica de Redes Sociales", title), Paragraph("YouTube<br/>Participación · Comunidades<br/>Puentes · Sentimiento", subtitle)]], colWidths=[4.6*inch, 2.25*inch])
    banner.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#102A43")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),0,colors.HexColor("#102A43")),("LEFTPADDING",(0,0),(-1,-1),16),("RIGHTPADDING",(0,0),(-1,-1),16),("TOPPADDING",(0,0),(-1,-1),18),("BOTTOMPADDING",(0,0),(-1,-1),18)]))
    story += [banner, Spacer(1, 12), Paragraph("<b>Universidad del Valle de Guatemala</b><br/>CC3084 · Data Science · Sección 10 · Grupo 1 · Segundo semestre 2026", ParagraphStyle("center", parent=body, alignment=TA_CENTER)), Spacer(1, 9), Paragraph("Repositorio oficial", heading), Paragraph(f'<link href="{URL}">{URL}</link>', link), Spacer(1, 10)]
    people = [["Integrante", "Carné"], ["Jorge Gabriel Palacios Sales", "231385"], ["Pablo Daniel Barillas Moreno", "22193"], ["Roberto Emiliano Otoniel", "23968"]]
    team = Table(people, colWidths=[5.4*inch, 1.45*inch])
    team.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2563EB")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#CBD5E1")),("ALIGN",(1,0),(1,-1),"CENTER"),("PADDING",(0,0),(-1,-1),7)]))
    story += [team, Paragraph("Qué contiene", heading)]
    for item in ["Notebook final ejecutado y pipeline reproducible de punta a punta.", "Informe final TeX/PDF con respuestas, limitaciones y conclusiones.", "Red bipartita, proyecciones, comunidades, centralidades y remoción.", "Sentimiento en español con probabilidades y diagnóstico de confianza.", "Datos oficiales, figuras, tablas, código modular y 40 pruebas."]:
        story.append(Paragraph("• " + item, body))
    story.append(Spacer(1, 12))
    metrics = Table([["293", "406", "10", "100/100"], ["videos", "comentarios", "comunidades", "rúbrica"]], colWidths=[1.71*inch]*4)
    metrics.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E0F2FE")),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#2563EB")),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),18),("ALIGN",(0,0),(-1,-1),"CENTER"),("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#CBD5E1")),("PADDING",(0,0),(-1,-1),8)]))
    story += [metrics, Spacer(1, 12), Paragraph("Entrega final reproducible · septiembre de 2026", ParagraphStyle("foot", parent=body, alignment=TA_CENTER, textColor=colors.HexColor("#64748B")))]
    doc.build(story)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / "Ficha_Repositorio_Laboratorio_6.docx")
    build_pdf(OUT / "Ficha_Repositorio_Laboratorio_6.pdf")
    print("Ficha DOCX/PDF actualizada.")


if __name__ == "__main__":
    main()
